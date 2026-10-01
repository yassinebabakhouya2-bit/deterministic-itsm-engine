"""Read exports: CSV (comma or semicolon, UTF-8 or Windows-1252), XLSX, JSON, JSONL.

Column names are matched case- and accent-insensitively, so "Numéro" matches
"numero".
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import unicodedata
from pathlib import Path

from .errors import InputError


def normalize_header(name: str) -> str:
    name = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode("ascii")
    return " ".join(name.replace("_", " ").lower().split())


def decode_bytes(data: bytes, encoding: str | None = None) -> tuple[str, str]:
    if encoding:
        return data.decode(encoding), encoding
    for candidate in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(candidate), candidate
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def guess_delimiter(header_line: str) -> str:
    counts = {sep: header_line.count(sep) for sep in (";", ",", "\t", "|")}
    best = max(counts, key=lambda sep: counts[sep])
    return best if counts[best] else ","


def cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return " | ".join(cell_to_str(item) for item in value)
    return str(value)


def _read_csv(path: Path, delimiter: str | None, encoding: str | None):
    text, used = decode_bytes(path.read_bytes(), encoding)
    first = next((line for line in text.splitlines() if line.strip()), "")
    sep = delimiter or guess_delimiter(first)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=sep)
    headers: list[str] | None = None
    rows: list[dict] = []
    for values in reader:
        if headers is None:
            if not any(v.strip() for v in values):
                continue
            headers = [v.strip().lstrip("﻿") or f"column_{i + 1}" for i, v in enumerate(values)]
            continue
        if not any(v.strip() for v in values):
            continue
        values = (values + [""] * len(headers))[: len(headers)]
        rows.append(dict(zip(headers, values)))
    return headers or [], rows, {"format": "csv", "encoding": used, "delimiter": sep}


def _read_xlsx(path: Path, sheet: str | None):
    try:
        import openpyxl
    except ImportError:
        raise InputError("reading .xlsx needs openpyxl (pip install openpyxl), or save the export as CSV") from None
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise InputError(f"sheet {sheet!r} not found; sheets: {', '.join(workbook.sheetnames)}")
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[0]
        rows_iter = worksheet.iter_rows(values_only=True)
        headers: list[str] = []
        for values in rows_iter:
            if any(v not in (None, "") for v in values):
                headers = [cell_to_str(v).strip() or f"column_{i + 1}" for i, v in enumerate(values)]
                break
        rows = []
        for values in rows_iter:
            cells = [cell_to_str(v) for v in values]
            if not any(c.strip() for c in cells):
                continue
            cells = (cells + [""] * len(headers))[: len(headers)]
            rows.append(dict(zip(headers, cells)))
        return headers, rows, {"format": "xlsx", "sheet": worksheet.title}
    finally:
        workbook.close()


def _read_json(path: Path):
    text, used = decode_bytes(path.read_bytes())
    if path.suffix.lower() == ".json":
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            raise InputError(f"{path}: invalid JSON ({exc.msg})") from None
        if isinstance(data, dict):
            data = next((v for v in data.values() if isinstance(v, list)), None)
        if not isinstance(data, list):
            raise InputError(f"{path}: expected a list of objects")
        objects = data
    else:
        objects = []
        for lineno, line in enumerate(text.splitlines(), 1):
            if not line.strip():
                continue
            try:
                objects.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise InputError(f"{path}:{lineno}: invalid JSON ({exc.msg})") from None
    headers: list[str] = []
    for obj in objects:
        if not isinstance(obj, dict):
            raise InputError(f"{path}: every record must be a JSON object")
        for key in obj:
            if key not in headers:
                headers.append(key)
    return headers, objects, {"format": "json", "encoding": used}


def read_rows(path: str | Path, delimiter: str | None = None, encoding: str | None = None, sheet: str | None = None):
    """Return (headers, rows as dicts, info) for a CSV, XLSX, JSON or JSONL file."""
    path = Path(path)
    if not path.is_file():
        raise InputError(f"file not found: {path}")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet)
    if suffix in (".json", ".jsonl", ".ndjson"):
        return _read_json(path)
    return _read_csv(path, delimiter, encoding)


def resolve_column(headers: list[str], wanted: str, role: str) -> str:
    if wanted in headers:
        return wanted
    by_norm: dict[str, str] = {}
    for header in headers:
        by_norm.setdefault(normalize_header(header), header)
    found = by_norm.get(normalize_header(wanted))
    if found is None:
        listing = ", ".join(repr(h) for h in headers) or "(none)"
        raise InputError(f"{role} column {wanted!r} not found. Columns in the file: {listing}")
    return found
