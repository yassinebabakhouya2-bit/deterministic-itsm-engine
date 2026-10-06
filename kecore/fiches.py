"""KB fiches: the input of the decomposition, and how they are read.

A fiche comes from a folder of documents (Markdown, text, HTML, and Word or
PDF when python-docx or pypdf is installed), from a JSONL file, or from a KB
export table (EasyVista, ServiceNow) through ``import_fiches``.

Its id follows kecore.ids: the ServiceNow number in the file name or title
(KB0012345) when there is one, otherwise the file name without its extension,
so that the engine, the search index and the scoreboard name fiches alike.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from .errors import InputError
from .ids import DEFAULT_FICHE_REGEX, compile_pattern, strip_document_path
from .tables import cell_to_str, decode_bytes, read_rows, resolve_column
from .segment import heading_key, keyword_role
from .text import clean_text

DOCUMENT_EXTENSIONS = frozenset({".md", ".markdown", ".txt", ".html", ".htm", ".docx", ".pdf"})
_HEADING_RE = re.compile(r"^\s{0,3}#{1,2}\s+(.+?)\s*#*\s*$")
_TITLE_LABEL_RE = re.compile(r"(?i)^\s*(?:titre|title|objet|sujet)\s*:\s*(.+)$")


@dataclass
class Fiche:
    fiche_id: str
    client: str
    title: str
    text: str
    source: str = ""
    meta: dict = field(default_factory=dict)

    @property
    def key(self) -> str:
        return f"{self.client}/{self.fiche_id}"

    def to_dict(self) -> dict:
        data = {"fiche_id": self.fiche_id, "client": self.client, "title": self.title, "text": self.text}
        if self.source:
            data["source"] = self.source
        if self.meta:
            data["meta"] = self.meta
        return data


def _read_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise InputError("reading .docx needs python-docx (pip install python-docx)") from None
    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            lines.append("")
            continue
        style = (paragraph.style.name if paragraph.style is not None else "").lower()
        properties = paragraph._p.pPr
        numbered = properties is not None and properties.numPr is not None
        if style.startswith(("heading", "titre")):
            lines.append("## " + text)
        elif numbered or "list" in style or "liste" in style:
            lines.append("- " + text)
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def _read_pdf(data: bytes) -> str:
    try:
        import pypdf
    except ImportError:
        raise InputError("reading .pdf needs pypdf (pip install pypdf)") from None
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def read_document_bytes(name: str, data: bytes) -> str:
    """Text of one document from its bytes: the same parsing wherever the bytes come from
    (a local folder, or a blob read by the Azure Function)."""
    suffix = Path(name).suffix.lower()
    if suffix == ".docx":
        return _read_docx(data)
    if suffix == ".pdf":
        return _read_pdf(data)
    return decode_bytes(data)[0]


def read_document(path: Path) -> str:
    path = Path(path)
    return read_document_bytes(path.name, path.read_bytes())


_EMPHASIS_RE = re.compile(r"(\*{1,3}|_{2,3}|`)")


def _clean_title(title: str) -> str:
    return " ".join(_EMPHASIS_RE.sub("", title.lstrip("#")).split())


def extract_title(text: str) -> str | None:
    for line in [line for line in text.split("\n") if line.strip()][:5]:
        for pattern in (_HEADING_RE, _TITLE_LABEL_RE):
            match = pattern.match(line)
            if match:
                return _clean_title(match.group(1)) or None
    first = next((line.strip() for line in text.split("\n") if line.strip()), "")
    if first and len(first) <= 150:
        return _clean_title(first) or None
    return None


def fiche_id_for(name: str, title: str, pattern) -> str:
    for candidate in (name, title):
        if pattern is not None and candidate:
            match = pattern.search(candidate)
            if match:
                return match.group(0)
    return strip_document_path(name)


def document_sort_key(relative: str) -> tuple[str, ...]:
    """Order in which documents are read: case-insensitive and separator-agnostic.

    It is the order Windows gave the reference runs (sorting WindowsPath objects compares
    lowercased parts), now identical on every OS, so a run on Linux in Azure reads the same
    documents in the same order. Order matters twice: duplicate fiche ids keep the first file,
    and the profile keeps the first spelling it meets for each heading.
    """
    return tuple(part.lower() for part in relative.replace("\\", "/").split("/"))


def fiches_from_documents(documents, client: str, id_regex: str | None = DEFAULT_FICHE_REGEX) -> tuple[list[Fiche], list[str]]:
    """Fiches from ``(relative_name, read)`` pairs, ``read()`` returning the document's bytes.

    The one place a document becomes a fiche, for a local folder (``load_folder``) and for blobs
    read in Azure alike: same order, same text extraction, same title and id rules.
    """
    pattern = compile_pattern(id_regex)
    fiches: list[Fiche] = []
    warnings: list[str] = []
    seen: dict[str, str] = {}
    for relative, read in sorted(documents, key=lambda item: document_sort_key(item[0])):
        name = relative.replace("\\", "/").rsplit("/", 1)[-1]
        try:
            text = clean_text(read_document_bytes(name, read()))
        except InputError as exc:
            warnings.append(f"{relative}: {exc}")
            continue
        except Exception as exc:  # a damaged document must not stop the whole KB
            warnings.append(f"{relative}: unreadable ({type(exc).__name__}: {exc})")
            continue
        if not text:
            warnings.append(f"{relative}: no text")
            continue
        title = extract_title(text)
        if not title or keyword_role(heading_key(title)) is not None:
            title = Path(name).stem  # a first line such as "Description" is a section, not a title
        fiche_id = fiche_id_for(name, title, pattern)
        if fiche_id in seen:
            warnings.append(f"{relative}: same fiche id {fiche_id} as {seen[fiche_id]}, ignored")
            continue
        seen[fiche_id] = relative
        fiches.append(Fiche(fiche_id=fiche_id, client=client, title=title, text=text, source=relative))
    return fiches, warnings


def load_folder(folder: str | Path, client: str, id_regex: str | None = DEFAULT_FICHE_REGEX) -> tuple[list[Fiche], list[str]]:
    folder = Path(folder)
    paths = [p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in DOCUMENT_EXTENSIONS]
    documents = [(str(p.relative_to(folder)), p.read_bytes) for p in paths]
    return fiches_from_documents(documents, client, id_regex)


def import_fiches(path: str | Path, client: str, body_cols: list[str], id_col: str, title_col: str | None = None,
                  sheet: str | None = None, delimiter: str | None = None, encoding: str | None = None,
                  id_regex: str | None = DEFAULT_FICHE_REGEX) -> tuple[list[Fiche], list[str]]:
    """Fiches from a KB export table. Several body columns become sections named after them."""
    headers, rows, _ = read_rows(path, delimiter, encoding, sheet)
    id_column = resolve_column(headers, id_col, "id")
    title_column = resolve_column(headers, title_col, "title") if title_col else None
    body_columns = [resolve_column(headers, c, "body") for c in body_cols]
    pattern = compile_pattern(id_regex)
    fiches: list[Fiche] = []
    warnings: list[str] = []
    seen: set[str] = set()
    for line, row in enumerate(rows, 2):
        raw_id = cell_to_str(row.get(id_column)).strip()
        title = clean_text(cell_to_str(row.get(title_column))) if title_column else ""
        parts = []
        for column in body_columns:
            body = clean_text(cell_to_str(row.get(column)))
            if body:
                parts.append(f"## {column}\n{body}" if len(body_columns) > 1 else body)
        if not raw_id or not parts:
            warnings.append(f"row {line}: no id or no text, skipped")
            continue
        fiche_id = fiche_id_for(raw_id, title, pattern)
        if fiche_id in seen:
            warnings.append(f"row {line}: duplicate fiche id {fiche_id}, skipped")
            continue
        seen.add(fiche_id)
        text = (f"# {title}\n\n" if title else "") + "\n\n".join(parts)
        fiches.append(Fiche(fiche_id=fiche_id, client=client, title=title or fiche_id, text=text, source=f"row {line}"))
    return fiches, warnings


def save_fiches(path: str | Path, fiches: list[Fiche]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        for fiche in fiches:
            handle.write(json.dumps(fiche.to_dict(), ensure_ascii=False) + "\n")


def load_jsonl(path: str | Path, client: str | None = None) -> list[Fiche]:
    fiches = []
    path = Path(path)
    with path.open(encoding="utf-8-sig") as handle:
        for lineno, line in enumerate(handle, 1):
            if not line.strip():
                continue
            try:
                data = json.loads(line)
                fiche = Fiche(
                    fiche_id=str(data["fiche_id"]),
                    client=str(data.get("client") or client or ""),
                    title=str(data.get("title") or data["fiche_id"]),
                    text=str(data["text"]),
                    source=str(data.get("source") or ""),
                    meta=data.get("meta") or {},
                )
            except (json.JSONDecodeError, KeyError, TypeError) as exc:
                raise InputError(f"{path}:{lineno}: not a fiche ({exc})") from None
            if not fiche.client:
                raise InputError(f"{path}:{lineno}: no client (pass --client)")
            if client and fiche.client != client:
                continue
            fiches.append(fiche)
    return fiches


def load_fiches(source: str | Path, client: str, id_regex: str | None = DEFAULT_FICHE_REGEX) -> tuple[list[Fiche], list[str]]:
    source = Path(source)
    if source.is_dir():
        return load_folder(source, client, id_regex)
    if source.suffix.lower() in (".jsonl", ".ndjson"):
        return load_jsonl(source, client), []
    if source.is_file():
        raise InputError(f"{source}: give a folder of fiches or a .jsonl file (use import-fiches for CSV and XLSX exports)")
    raise InputError(f"not found: {source}")
