import unittest

from kecore.errors import InputError
from kecore.fiches import extract_title, import_fiches, load_fiches, load_folder, load_jsonl, save_fiches

from .helpers import DEMO_KB, TempDirTestCase

try:
    import docx
except ImportError:  # optional dependency
    docx = None


class FolderTest(TempDirTestCase):
    def test_demo_kb(self):
        fiches, warnings = load_folder(DEMO_KB, "clienta")
        self.assertEqual(warnings, [])
        ids = [f.fiche_id for f in fiches]
        self.assertEqual(sorted(ids), ["KB0010001", "KB0010002", "KB0010003", "KB0010005", "KB0010006", "KB0010020"])
        by_id = {f.fiche_id: f for f in fiches}
        self.assertEqual(by_id["KB0010003"].title, "KB0010003 - Imprimante en bourrage", "a section name is not a title")
        self.assertEqual(by_id["KB0010002"].title, "VPN - déconnexions fréquentes")
        self.assertIn("- Éteindre l'imprimante.", by_id["KB0010003"].text)
        self.assertTrue(all(f.client == "clienta" for f in fiches))

    def test_ids_without_number_and_duplicates(self):
        self.write("kb/Réinitialiser le mot de passe.md", "# Réinitialiser le mot de passe\nCliquez sur Oublié.")
        self.write("kb/KB0010009.txt", "Fiche A")
        self.write("kb/old/KB0010009 - ancienne version.md", "Fiche B")
        self.write("kb/vide.md", "   ")
        fiches, warnings = load_folder(self.path("kb"), "c")
        self.assertEqual(sorted(f.fiche_id for f in fiches), ["KB0010009", "Réinitialiser le mot de passe"])
        self.assertEqual(len(warnings), 2)
        self.assertTrue(any("same fiche id KB0010009" in w for w in warnings))

    def test_jsonl_round_trip_and_dispatch(self):
        fiches, _ = load_folder(DEMO_KB, "clienta")
        path = self.path("fiches.jsonl")
        save_fiches(path, fiches)
        self.assertEqual([f.to_dict() for f in load_jsonl(path)], [f.to_dict() for f in fiches])
        self.assertEqual(len(load_fiches(path, "clienta")[0]), 6)
        self.assertEqual(load_fiches(path, "clientb")[0], [])
        with self.assertRaisesRegex(InputError, "import-fiches"):
            load_fiches(self.write("x.csv", "a;b\n"), "c")


class ImportTest(TempDirTestCase):
    def test_export_with_two_text_columns(self):
        export = self.write(
            "kb.csv",
            "Référence;Titre;Description;Solution\n"
            "KB0010030;Écran noir;L'écran reste noir.;<ol><li>Rebranchez le câble.</li><li>Redémarrez.</li></ol>\n"
            "KB0010031;Vide;;\n",
            encoding="cp1252",
        )
        fiches, warnings = import_fiches(export, "client-s", ["description", "solution"], id_col="reference", title_col="titre")
        self.assertEqual([f.fiche_id for f in fiches], ["KB0010030"])
        self.assertEqual(
            fiches[0].text,
            "# Écran noir\n\n## Description\nL'écran reste noir.\n\n## Solution\n- Rebranchez le câble.\n\n- Redémarrez.",
        )
        self.assertEqual(len(warnings), 1)


class TitleTest(unittest.TestCase):
    def test_title(self):
        self.assertEqual(extract_title("\n# **Outlook** bloqué\ntexte"), "Outlook bloqué")
        self.assertEqual(extract_title("Titre : VPN coupé\nProblème : x"), "VPN coupé")
        self.assertEqual(extract_title("Première ligne\nSuite"), "Première ligne")


@unittest.skipIf(docx is None, "python-docx not installed")
class DocxTest(TempDirTestCase):
    def test_word_headings_and_lists(self):
        document = docx.Document()
        document.add_heading("Résolution", level=2)
        document.add_paragraph("Fermez Outlook.", style="List Number")
        document.add_paragraph("Relancez Outlook.", style="List Number")
        folder = self.path("kb")
        folder.mkdir()
        document.save(folder / "KB0010040 - Outlook.docx")
        fiches, warnings = load_folder(folder, "c")
        self.assertEqual(warnings, [])
        self.assertEqual(fiches[0].text, "## Résolution\n- Fermez Outlook.\n- Relancez Outlook.")


if __name__ == "__main__":
    unittest.main()
