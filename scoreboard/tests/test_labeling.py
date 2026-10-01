import csv
import unittest

from scoreboard.dataset import DatasetError, Ticket
from scoreboard.engines.fixture import FixtureEngine
from scoreboard.labeling import apply_labels, format_candidate, parse_candidate, prepare_labels

from .helpers import TempDirTestCase

try:
    import openpyxl
except ImportError:
    openpyxl = None


def sample_tickets():
    return [
        Ticket("0042", "clienta", "Outlook ne démarre plus"),
        Ticket("T-2", "clienta", "Demande d'un second écran"),
        Ticket("T-3", "clienta", "Imprimante en bourrage"),
        Ticket("T-4", "clienta", "Déjà étiqueté", expected=["KB9"]),
    ]


def sample_engine():
    return FixtureEngine(
        {},
        candidates={
            "clienta/0042": [{"id": "KB0010001", "title": "Outlook gelé au démarrage"}, {"id": "KB0010009", "title": "Réparer Office"}],
            "clienta/T-3": ["KB0010003"],
        },
    )


class CandidateFormatTest(unittest.TestCase):
    def test_round_trip(self):
        cell = format_candidate("KB1", "VPN — accès   distant")
        self.assertEqual(cell, "KB1 — VPN — accès distant")
        self.assertEqual(parse_candidate(cell), "KB1")
        self.assertEqual(format_candidate("KB2", ""), "KB2")


class CsvSheetTest(TempDirTestCase):
    def fill(self, path, values, encoding="utf-8-sig"):
        """Simulate Excel: rewrite the sheet with the 'expected' cells filled."""
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle, delimiter=";"))
        column = rows[0].index("expected")
        for row in rows[1:]:
            if row[0] in values:
                row[column] = values[row[0]]
        with path.open("w", encoding=encoding, newline="") as handle:
            csv.writer(handle, delimiter=";").writerows(rows)

    def test_prepare_then_apply(self):
        tickets = sample_tickets()
        sheet = self.path("labels.csv")
        report = prepare_labels(tickets, sheet, engine=sample_engine(), k=3)
        self.assertEqual((report.rows, report.with_candidates), (3, 2))
        raw = sheet.read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "French Excel needs the UTF-8 BOM")
        header = raw.decode("utf-8-sig").splitlines()[0]
        self.assertEqual(header, "ticket_id;client;text;cand_1;cand_2;cand_3;expected;notes")

        # Excel turns 0042 into 42 and the file may come back in Windows-1252.
        self.fill(sheet, {"0042": "2", "T-2": "aucune", "T-3": "KB0010003 | KB0010013"}, encoding="cp1252")
        content = sheet.read_text(encoding="cp1252").replace("\n0042;", "\n42;")
        sheet.write_text(content, encoding="cp1252")

        updated, applied = apply_labels(sheet, tickets)
        expected = {t.ticket_id: t.expected for t in updated}
        self.assertEqual(expected["0042"], ["KB0010009"])
        self.assertEqual(expected["T-2"], [])
        self.assertEqual(expected["T-3"], ["KB0010003", "KB0010013"])
        self.assertEqual(expected["T-4"], ["KB9"])
        self.assertEqual((applied.labeled, applied.none, applied.empty), (2, 1, 0))
        self.assertIsNone(tickets[0].expected, "the input tickets are not modified")

    def test_errors(self):
        tickets = sample_tickets()
        sheet = self.path("labels.csv")
        prepare_labels(tickets, sheet, engine=sample_engine(), k=3)
        self.fill(sheet, {"T-3": "1|none"})
        with self.assertRaisesRegex(DatasetError, "cannot be combined"):
            apply_labels(sheet, tickets)
        prepare_labels(tickets, sheet, engine=sample_engine(), k=3)
        self.fill(sheet, {"T-3": "2"})
        with self.assertRaisesRegex(DatasetError, "candidate 2 is empty"):
            apply_labels(sheet, tickets)

    def test_unknown_rows_are_reported(self):
        sheet = self.write("labels.csv", "ticket_id;client;expected\nZZ;clienta;KB1\n")
        _, report = apply_labels(sheet, sample_tickets())
        self.assertEqual(report.unknown, ["clienta/ZZ"])

    def test_include_labeled(self):
        sheet = self.path("all.csv")
        report = prepare_labels(sample_tickets(), sheet, include_labeled=True)
        self.assertEqual(report.rows, 4)
        self.assertIn("KB9", sheet.read_text(encoding="utf-8-sig"))


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class XlsxSheetTest(TempDirTestCase):
    def test_xlsx_round_trip_keeps_ids_as_text(self):
        tickets = sample_tickets()
        sheet = self.path("labels.xlsx")
        prepare_labels(tickets, sheet, engine=sample_engine(), k=2)
        workbook = openpyxl.load_workbook(sheet)
        ws = workbook.active
        header = [c.value for c in ws[1]]
        self.assertEqual(header, ["ticket_id", "client", "text", "cand_1", "cand_2", "expected", "notes"])
        self.assertEqual(ws["A2"].value, "0042")
        ws["F2"] = 1
        ws["F3"] = "none"
        workbook.save(sheet)
        updated, report = apply_labels(sheet, tickets)
        self.assertEqual(updated[0].expected, ["KB0010001"])
        self.assertEqual(updated[1].expected, [])
        self.assertEqual(report.labeled + report.none, 2)


if __name__ == "__main__":
    unittest.main()
